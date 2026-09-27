"""PPTX adapter: each slide becomes a 'page'; slide titles become headings."""
from __future__ import annotations

from pathlib import Path

from backend.ingestion.base import DocumentAdapter, ExtractionError, clean_text, register_adapter
from backend.schemas import Block, ContentType, ExtractedDocument


@register_adapter
class PPTXAdapter(DocumentAdapter):
    extensions = (".pptx",)
    format_name = "pptx"

    def extract(self, path: Path, document_id: str, work_dir: Path) -> ExtractedDocument:
        from pptx import Presentation
        from pptx.enum.shapes import MSO_SHAPE_TYPE

        try:
            prs = Presentation(str(path))
        except Exception as exc:
            raise ExtractionError("The PowerPoint file could not be opened. It may be corrupted.") from exc

        out = ExtractedDocument(document_id=document_id, filename=path.name, format="pptx", page_count=len(prs.slides))
        table_no = figure_no = 0

        def walk(shapes):
            for sh in shapes:
                if sh.shape_type == MSO_SHAPE_TYPE.GROUP:
                    yield from walk(sh.shapes)
                else:
                    yield sh

        for sno, slide in enumerate(prs.slides, start=1):
            title_shape = slide.shapes.title
            if title_shape is not None and title_shape.has_text_frame and clean_text(title_shape.text_frame.text):
                out.blocks.append(Block(type=ContentType.heading, text=clean_text(title_shape.text_frame.text),
                                        page=sno, level=1 if sno == 1 else 2))
            for sh in walk(slide.shapes):
                if title_shape is not None and sh.shape_id == title_shape.shape_id:
                    continue
                if getattr(sh, "has_table", False) and sh.has_table:
                    rows = [[clean_text(c.text) for c in r.cells] for r in sh.table.rows]
                    rows = [r for r in rows if any(r)]
                    if len(rows) >= 2:
                        table_no += 1
                        out.blocks.append(Block(type=ContentType.table, text="\n".join(" | ".join(r) for r in rows),
                                                page=sno, table=rows, table_index=table_no))
                elif sh.shape_type == MSO_SHAPE_TYPE.PICTURE:
                    try:
                        if sh.width and sh.width > prs.slide_width * 0.25:
                            figure_no += 1
                            rel = f"figures/figure_{figure_no:02d}.{sh.image.ext}"
                            (work_dir / "figures").mkdir(parents=True, exist_ok=True)
                            (work_dir / rel).write_bytes(sh.image.blob)
                            out.blocks.append(Block(type=ContentType.figure, text=f"Image on slide {sno}", page=sno,
                                                    figure_index=figure_no, image_path=rel))
                    except Exception:
                        pass
                elif sh.has_text_frame:
                    for para in sh.text_frame.paragraphs:
                        text = clean_text("".join(r.text for r in para.runs))
                        if text:
                            kind = ContentType.list_item if para.level > 0 or len(sh.text_frame.paragraphs) > 1 else ContentType.paragraph
                            out.blocks.append(Block(type=kind, text=text, page=sno))
            if slide.has_notes_slide:
                notes = clean_text(slide.notes_slide.notes_text_frame.text)
                if notes:
                    out.blocks.append(Block(type=ContentType.paragraph, text=notes, page=sno))

        if not out.blocks:
            raise ExtractionError("The PowerPoint file does not contain readable text.")
        out.title = next((b.text for b in out.blocks if b.type == ContentType.heading), path.stem)
        cp = prs.core_properties
        out.metadata = {k: v for k, v in {"title": cp.title, "author": cp.author}.items() if v}
        return out
