"""PPTX backend (python-pptx): Scene -> native, editable PowerPoint elements.

If the brand has a reference template, it is used as the base presentation so its slide
masters, backgrounds and fixed artwork carry through; slides are built on its blank layout.
"""
from __future__ import annotations

import copy
from pathlib import Path
from typing import Optional

from lxml import etree
from PIL import Image
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION, XL_LEGEND_POSITION
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

from backend.branding.theme import Theme
from backend.rendering.charts import fmt_value
from backend.rendering.scene import Chart, Line, Picture, Scene, Shape, Table, TextBox, fit_text
from backend.utils.logging import get_logger

log = get_logger(__name__)

SHAPES = {"rect": MSO_SHAPE.RECTANGLE, "round_rect": MSO_SHAPE.ROUNDED_RECTANGLE, "ellipse": MSO_SHAPE.OVAL,
          "chevron": MSO_SHAPE.CHEVRON, "pentagon": MSO_SHAPE.PENTAGON, "right_arrow": MSO_SHAPE.RIGHT_ARROW,
          "triangle": MSO_SHAPE.ISOSCELES_TRIANGLE, "diamond": MSO_SHAPE.DIAMOND}
ALIGN = {"left": PP_ALIGN.LEFT, "center": PP_ALIGN.CENTER, "right": PP_ALIGN.RIGHT}
ANCHOR = {"top": MSO_ANCHOR.TOP, "middle": MSO_ANCHOR.MIDDLE, "bottom": MSO_ANCHOR.BOTTOM}


def rgb(h: str) -> RGBColor:
    return RGBColor.from_string(h.lstrip("#").upper())


def _clear_slides(prs) -> None:
    sld_ids = prs.slides._sldIdLst
    for sld in list(sld_ids):
        prs.part.drop_rel(sld.get(qn("r:id")))
        sld_ids.remove(sld)


def _blank_layout(prs, preferred: Optional[str]):
    layouts = list(prs.slide_layouts)
    if preferred:
        for l in layouts:
            if l.name == preferred:
                return l
    return min(layouts, key=lambda l: (len(l.placeholders), 0 if "blank" in l.name.lower() else 1))


class PPTXGenerator:
    def __init__(self, theme: Theme):
        self.theme = theme
        self.using_template = bool(theme.template_path)

    def new_presentation(self):
        t = self.theme
        if self.using_template:
            try:
                prs = Presentation(str(t.template_path))
                _clear_slides(prs)
                return prs
            except Exception as exc:
                log.warning("Template could not be used; generating without it", error=type(exc).__name__)
                self.using_template = False
        prs = Presentation()
        prs.slide_width, prs.slide_height = Inches(t.slide_w), Inches(t.slide_h)
        return prs

    def build(self, scenes: list[Scene], out_path: Path, title: str = "") -> Path:
        prs = self.new_presentation()
        # scale scenes if a template's slide size differs from the theme size
        sx = prs.slide_width / Inches(self.theme.slide_w)
        sy = prs.slide_height / Inches(self.theme.slide_h)
        self._scale = min(sx, sy)
        layout = _blank_layout(prs, self.theme.template_blank_layout)
        for scene in scenes:
            slide = prs.slides.add_slide(layout)
            for ph in list(slide.placeholders):
                ph._element.getparent().remove(ph._element)
            if not (self.using_template and (scene.use_template_background or scene.background == self.theme.c("background"))):
                fill = slide.background.fill
                fill.solid()
                fill.fore_color.rgb = rgb(scene.background)
            for el in scene.elements:
                try:
                    self._add(slide, el)
                except Exception as exc:  # one bad element must not break the deck
                    log.warning("Element skipped", slide=scene.name, element=type(el).__name__, error=type(exc).__name__)
            if scene.notes:
                slide.notes_slide.notes_text_frame.text = scene.notes
        cp = prs.core_properties
        cp.title = title[:250]
        cp.author = ""
        cp.comments = "Generated offline by Content Studio"
        out_path.parent.mkdir(parents=True, exist_ok=True)
        prs.save(str(out_path))
        log.info("PPT rendering completed", slides=len(scenes))
        return out_path

    # -------------------------------------------------------------- helpers
    def _e(self, inches: float) -> Emu:
        return Emu(int(Inches(inches) * self._scale))

    def _add(self, slide, el) -> None:
        if isinstance(el, TextBox):
            self._text(slide, el)
        elif isinstance(el, Shape):
            self._shape(slide, el)
        elif isinstance(el, Line):
            self._line(slide, el)
        elif isinstance(el, Picture):
            self._picture(slide, el)
        elif isinstance(el, Table):
            self._table(slide, el)
        elif isinstance(el, Chart):
            self._chart(slide, el)

    def _text(self, slide, tb: TextBox, frame=None) -> None:
        if not tb.fitted_size:
            fit_text(tb)
        if frame is None:
            box = slide.shapes.add_textbox(self._e(tb.x), self._e(tb.y), self._e(tb.w), self._e(tb.h))
            if tb.name:
                box.name = tb.name
            frame = box.text_frame
        tf = frame
        tf.word_wrap = True
        tf.auto_size = MSO_AUTO_SIZE.NONE
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        tf.vertical_anchor = ANCHOR[tb.valign]
        size = tb.fitted_size * self._scale
        for i, p in enumerate(tb.paras):
            para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            para.alignment = ALIGN[tb.align]
            para.line_spacing = 1.0
            ps = size * p.size_scale
            if i < len(tb.paras) - 1:
                para.space_after = Pt(ps * p.space_after)
            run = para.add_run()
            run.text = p.text
            f = run.font
            f.name = tb.font
            f.size = Pt(round(ps, 1))
            f.bold = p.bold if p.bold is not None else tb.bold
            f.italic = tb.italic
            f.color.rgb = rgb(p.color or tb.color)
            if p.bullet:
                self._bullet(para, p.bullet, tb.bullet_color or tb.color, ps)

    @staticmethod
    def _bullet(para, char: str, color: str, size_pt: float) -> None:
        pPr = para._p.get_or_add_pPr()
        indent = int(Pt(size_pt * 1.3))
        pPr.set("marL", str(indent))
        pPr.set("indent", str(-indent))
        for tag in ("a:buClr", "a:buFont", "a:buChar", "a:buAutoNum", "a:buNone"):
            for old in pPr.findall(qn(tag)):
                pPr.remove(old)
        buClr = etree.SubElement(pPr, qn("a:buClr"))
        clr = etree.SubElement(buClr, qn("a:srgbClr"))
        clr.set("val", color.lstrip("#").upper())
        if char.endswith(".") and char[:-1].isdigit():
            num = etree.SubElement(pPr, qn("a:buAutoNum"))
            num.set("type", "arabicPeriod")
            num.set("startAt", char[:-1])
        else:
            bf = etree.SubElement(pPr, qn("a:buFont"))
            bf.set("typeface", "Arial")
            bc = etree.SubElement(pPr, qn("a:buChar"))
            bc.set("char", char)

    def _freeform(self, slide, sh: Shape):
        from dataclasses import replace

        from backend.rendering.images.raster import shape_polygon

        pts = shape_polygon(replace(sh, rotation=0.0), 1.0)  # inches; rotation applied as a shape property
        fb = slide.shapes.build_freeform(self._e(pts[0][0]), self._e(pts[0][1]), scale=1.0)
        fb.add_line_segments([(self._e(x), self._e(y)) for x, y in pts[1:]], close=True)
        return fb.convert_to_shape()

    def _shape(self, slide, sh: Shape) -> None:
        if sh.kind in ("right_triangle", "quarter_circle"):
            shp = self._freeform(slide, sh)
        else:
            kind = MSO_SHAPE.ROUNDED_RECTANGLE if sh.kind == "pill" else SHAPES[sh.kind]
            shp = slide.shapes.add_shape(kind, self._e(sh.x), self._e(sh.y), self._e(sh.w), self._e(sh.h))
            if sh.kind == "pill":
                shp.adjustments[0] = 0.5
        if sh.rotation % 360:
            shp.rotation = sh.rotation
        if sh.name:
            shp.name = sh.name
        if sh.fill:
            shp.fill.solid()
            shp.fill.fore_color.rgb = rgb(sh.fill)
            if sh.opacity < 0.999:
                clr = shp.fill._xPr.find(qn("a:solidFill")).find(qn("a:srgbClr"))
                alpha = etree.SubElement(clr, qn("a:alpha"))
                alpha.set("val", str(int(max(0.0, sh.opacity) * 100000)))
        else:
            shp.fill.background()
        if sh.line and sh.line_pt:
            shp.line.color.rgb = rgb(sh.line)
            shp.line.width = Pt(sh.line_pt)
        else:
            shp.line.fill.background()
        if sh.kind == "round_rect" and min(sh.w, sh.h) > 0:
            shp.adjustments[0] = max(0.0, min(0.5, sh.radius / min(sh.w, sh.h)))
        if sh.kind in ("chevron", "pentagon"):
            tip = min(sh.h * 0.35, sh.w * 0.25)
            shp.adjustments[0] = max(0.0, min(1.0, tip / sh.h))
        shp.shadow.inherit = False
        # no default text styling on decorative shapes
        shp.text_frame.text = ""

    def _line(self, slide, ln: Line) -> None:
        c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, self._e(ln.x1), self._e(ln.y1), self._e(ln.x2), self._e(ln.y2))
        c.line.color.rgb = rgb(ln.color)
        c.line.width = Pt(ln.width_pt)
        if ln.arrow_end:
            lnel = c.line._get_or_add_ln()
            tail = etree.SubElement(lnel, qn("a:tailEnd"))
            tail.set("type", "triangle")
            tail.set("w", "med")
            tail.set("len", "med")

    def _picture(self, slide, pic: Picture) -> None:
        if pic.name == "logo" and self.using_template and "template_logo" in Path(pic.path).name:
            return  # the template master already carries this logo
        with Image.open(pic.path) as im:
            iw, ih = im.size
        if pic.fit == "cover":  # native PowerPoint crop keeps the full image editable
            p = slide.shapes.add_picture(pic.path, self._e(pic.x), self._e(pic.y), self._e(pic.w), self._e(pic.h))
            box_ar, img_ar = pic.w / pic.h, iw / ih
            if img_ar > box_ar:
                c = (1 - box_ar / img_ar) / 2
                p.crop_left = p.crop_right = c
            else:
                c = (1 - img_ar / box_ar) / 2
                p.crop_top = p.crop_bottom = c
        else:
            r = min(pic.w / iw, pic.h / ih)
            w, h = iw * r, ih * r
            p = slide.shapes.add_picture(pic.path, self._e(pic.x + (pic.w - w) / 2), self._e(pic.y + (pic.h - h) / 2),
                                         self._e(w), self._e(h))
        if pic.name:
            p.name = pic.name

    def _table(self, slide, tb: Table) -> None:
        from backend.rendering.images.raster import table_cells

        cells = table_cells(tb)
        ncols = max(len(tb.headers), max((len(r) for r in tb.rows), default=0))
        nrows = len(tb.rows) + 1
        gf = slide.shapes.add_table(nrows, ncols, self._e(tb.x), self._e(tb.y), self._e(tb.w), self._e(tb.h))
        gf.name = "table"
        table = gf.table
        tblPr = table._tbl.tblPr
        for attr in ("firstRow", "bandRow"):
            tblPr.set(attr, "0")
        style = tblPr.find(qn("a:tableStyleId"))
        if style is not None:
            style.text = "{5940675A-B579-460E-94D1-54222C63F5DA}"  # 'No Style, Table Grid' -> we paint cells ourselves
        widths = [c[3] for c in cells[:ncols]]
        for j, w in enumerate(widths):
            table.columns[j].width = self._e(w)
        for i in range(nrows):
            table.rows[i].height = self._e(tb.h / nrows)
        for k, (_, _, _, _, _, fill, box) in enumerate(cells):
            i, j = divmod(k, ncols)
            cell = table.cell(i, j)
            cell.margin_left = cell.margin_right = self._e(0.1)
            cell.margin_top = cell.margin_bottom = self._e(0.04)
            if fill:
                cell.fill.solid()
                cell.fill.fore_color.rgb = rgb(fill)
            else:
                cell.fill.background()
            cell.text_frame.text = ""
            self._text(slide, copy.copy(box), frame=cell.text_frame)

    def _chart(self, slide, ch: Chart) -> None:
        d = ch.data
        cd = CategoryChartData()
        cd.categories = d.categories
        for s in d.series[: (1 if d.chart_type == "pie" else 6)]:
            cd.add_series(s.name, s.values)
        kind = {"column": XL_CHART_TYPE.COLUMN_CLUSTERED, "bar": XL_CHART_TYPE.BAR_CLUSTERED,
                "line": XL_CHART_TYPE.LINE_MARKERS, "pie": XL_CHART_TYPE.PIE}[d.chart_type]
        gf = slide.shapes.add_chart(kind, self._e(ch.x), self._e(ch.y), self._e(ch.w), self._e(ch.h), cd)
        gf.name = "chart"
        chart = gf.chart
        chart.font.name = ch.font
        chart.font.size = Pt(12 * self._scale)
        chart.font.color.rgb = rgb(ch.text_color)
        chart.has_legend = len(d.series) > 1 or d.chart_type == "pie"
        if chart.has_legend:
            chart.legend.position = XL_LEGEND_POSITION.TOP
            chart.legend.include_in_layout = False
        plot = chart.plots[0]
        if d.chart_type == "pie":
            for i, point in enumerate(plot.series[0].points):
                point.format.fill.solid()
                point.format.fill.fore_color.rgb = rgb(ch.palette[i % len(ch.palette)])
        else:
            if hasattr(plot, "gap_width"):
                plot.gap_width = 60
            for i, s in enumerate(plot.series):
                col = rgb(ch.palette[i % len(ch.palette)])
                if d.chart_type == "line":
                    s.format.line.color.rgb = col
                    s.format.line.width = Pt(2.5)
                else:
                    s.format.fill.solid()
                    s.format.fill.fore_color.rgb = col
            va = chart.value_axis
            va.has_major_gridlines = ch.gridlines
            if ch.gridlines:
                va.major_gridlines.format.line.color.rgb = rgb(ch.grid_color)
            va.format.line.fill.background()
            chart.category_axis.format.line.color.rgb = rgb(ch.grid_color)
            chart.category_axis.tick_labels.font.size = Pt(11 * self._scale)
            va.tick_labels.font.size = Pt(10 * self._scale)
            if d.chart_type == "bar":
                chart.category_axis.reverse_order = True
        if ch.data_labels:
            plot.has_data_labels = True
            dl = plot.data_labels
            dl.font.size = Pt(10 * self._scale)
            dl.font.color.rgb = rgb(ch.text_color)
            if d.chart_type in ("column", "bar"):
                dl.position = XL_LABEL_POSITION.OUTSIDE_END
            if d.unit == "%" and d.chart_type != "pie":
                dl.number_format = '0.0"%"' if any(not float(v).is_integer() for s in d.series for v in s.values) else '0"%"'
                dl.number_format_is_linked = False


__all__ = ["PPTXGenerator", "fmt_value"]
