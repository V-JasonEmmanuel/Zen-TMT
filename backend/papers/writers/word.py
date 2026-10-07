"""Word (.docx) in the target format: page size and margins, one or two columns, the template's
fonts and sizes, numbered headings, caption conventions, booktabs-style tables, equations (Word's
own equations kept as OMML; LaTeX equations typeset), and the reference list with labels or hanging
indents. Built-in styles (Title, Heading 1-3, Caption) are used so the navigation pane works.
"""
from __future__ import annotations

import copy
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

from backend.papers.cite_render import Rendered
from backend.papers.formats import Format
from backend.papers.math_render import as_text, render as render_math
from backend.papers.model import Block, Inline, Paper
from backend.papers.writers.common import PLACEHOLDER, Layout, ack_title, arrange, caption_label, statement_title

PAGE = {"A4": (8.27, 11.69), "Letter": (8.5, 11.0)}
ALIGN = {"left": WD_ALIGN_PARAGRAPH.LEFT, "center": WD_ALIGN_PARAGRAPH.CENTER, "justify": WD_ALIGN_PARAGRAPH.JUSTIFY}


class WordWriter:
    def __init__(self, paper: Paper, fmt: Format, rend: Rendered, paper_dir: Path, style: str):
        self.p, self.f, self.r, self.dir, self.style = paper, fmt, rend, paper_dir, style
        self.doc = Document()
        self.lay: Layout = arrange(paper, fmt)
        self.math_dir = paper_dir / "outputs" / "_math"
        self.math_as_text = 0
        pw, ph = PAGE[fmt.page]
        top, right, bottom, left = fmt.margins_in
        self.col_w = (pw - left - right - (fmt.column_gap_in if fmt.columns == 2 else 0)) / fmt.columns
        self.text_w = pw - left - right

    # ------------------------------------------------------------------ setup
    def _setup(self):
        f = self.f
        s = self.doc.sections[0]
        s.page_width, s.page_height = Inches(PAGE[f.page][0]), Inches(PAGE[f.page][1])
        s.top_margin, s.right_margin, s.bottom_margin, s.left_margin = (Inches(x) for x in f.margins_in)
        st = self.doc.styles
        normal = st["Normal"]
        normal.font.name = f.font
        normal.element.rPr.rFonts.set(qn("w:eastAsia"), f.font)
        normal.font.size = Pt(f.body_pt)
        pf = normal.paragraph_format
        pf.space_before, pf.space_after = Pt(0), Pt(0)
        pf.line_spacing = f.line_spacing
        for name, size, bold, italic in (("Title", f.title_pt, f.title_bold, False), ("Heading 1", f.h1_pt, True, False),
                                         ("Heading 2", f.h2_pt, not f.h2_italic, f.h2_italic), ("Heading 3", f.body_pt, f.id != "ieee", True),
                                         ("Caption", f.caption_pt, False, False)):
            x = st[name]
            x.font.name, x.font.size, x.font.bold, x.font.italic = f.font, Pt(size), bold, italic
            x.font.color.rgb = RGBColor(0, 0, 0)
            rpr = x.element.get_or_add_rPr()
            rf = rpr.find(qn("w:rFonts"))
            if rf is None:
                rf = OxmlElement("w:rFonts")
                rpr.append(rf)
            for a in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
                rf.set(qn(a), f.font)
            for a in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
                if rf.get(qn(a)) is not None:
                    del rf.attrib[qn(a)]
            x.paragraph_format.space_before = Pt(size * 0.9 if name.startswith("Heading") else 0)
            x.paragraph_format.space_after = Pt(size * 0.4 if name.startswith("Heading") else 0)
            x.paragraph_format.keep_with_next = name.startswith("Heading")
            if name == "Title":  # remove the default bottom border
                ppr = x.element.get_or_add_pPr()
                for b in ppr.findall(qn("w:pBdr")):
                    ppr.remove(b)

    def _para(self, style: str | None = None, align: str | None = None, indent: float = 0.0, space_after: float = 0.0):
        p = self.doc.add_paragraph(style=style) if style else self.doc.add_paragraph()
        if align:
            p.alignment = ALIGN[align]
        elif self.f.justify and not style:
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        if indent:
            p.paragraph_format.first_line_indent = Inches(indent)
        if space_after:
            p.paragraph_format.space_after = Pt(space_after)
        return p

    def _runs(self, p, runs: list[Inline], size: float | None = None, bold: bool | None = None, italic: bool | None = None):
        for r in runs:
            if r.cite:
                txt = self.r.cite_text(list(r.cite), r.t == "narrative", r.raw_cite)
                run = p.add_run(txt)
            elif r.xref:
                from backend.papers.crossrefs import render as xref_text

                run = p.add_run(xref_text(self.f, self.lay, r.xref, r.t))
            elif r.math:
                if not self._inline_math(p, r.math, size or self.f.body_pt):
                    run = p.add_run(as_text(r.math))
                    run.italic = True
                    self.math_as_text += 1
                continue
            else:
                run = p.add_run(r.t)
                run.bold = r.b or None
                run.italic = r.i or None
                if r.sup:
                    run.font.superscript = True
                if r.sub:
                    run.font.subscript = True
            if size:
                run.font.size = Pt(size)
            if bold is not None:
                run.bold = bold or run.bold
            if italic is not None:
                run.italic = italic or run.italic

    def _inline_math(self, p, latex: str, size: float) -> bool:
        res = render_math(latex, self.math_dir, size_pt=size, inline=True)
        if not res:
            return False
        path, w, h = res
        p.add_run().add_picture(str(path), height=Pt(h))
        return True

    # ------------------------------------------------------------------ front matter
    def _front(self):
        f, p = self.f, self.p
        t = self._para("Title", align=f.title_align, space_after=10)
        if f.title_page:
            t.paragraph_format.space_before = Pt(150)
        t.add_run(p.title)
        if p.authors:
            ap = self._para(align=f.title_align, space_after=4)
            for i, a in enumerate(p.authors):
                if i:
                    ap.add_run(", " if i < len(p.authors) - 1 else (" and " if f.id in ("ieee", "acm", "apa7") else ", "))
                ap.add_run(a.name).font.size = Pt(f.body_pt + 1)
                marks = ",".join(str(x + 1) for x in a.affiliations) if len(p.affiliations) > 1 else ""
                if a.corresponding:
                    marks = (marks + "*") if marks else "*"
                if marks:
                    m = ap.add_run(marks)
                    m.font.superscript = True
        for i, aff in enumerate(p.affiliations):
            ap = self._para(align=f.title_align)
            if len(p.affiliations) > 1:
                m = ap.add_run(str(i + 1))
                m.font.superscript = True
            r = ap.add_run(aff)
            r.font.size = Pt(max(8, f.body_pt - 1))
            r.italic = f.id in ("ieee", "springer_lncs")
        emails = [a.email for a in p.authors if a.email]
        if emails:
            ep = self._para(align=f.title_align, space_after=10)
            r = ep.add_run(", ".join(emails))
            r.font.size = Pt(max(8, f.body_pt - 1))
        else:
            self._para(space_after=8)
        if f.title_page:  # APA: the abstract starts on its own page
            self.doc.paragraphs[-1].add_run().add_break(WD_BREAK.PAGE)
        if p.abstract:
            if f.abstract_runin:
                ab = self._para(align="justify", space_after=6)
                lab = ab.add_run(f.abstract_label + ("—" if f.id == "ieee" else " "))
                lab.bold = "bold" in f.abstract_label_style
                lab.italic = "italic" in f.abstract_label_style
                self._runs(ab, p.abstract, size=f.body_pt - (1 if f.id == "ieee" else 0), bold=f.id == "ieee" or None)
            else:
                h = self._para(align="left" if f.abstract_label_style != "heading" or f.id != "apa7" else "center")
                hr = h.add_run(f.abstract_label)
                hr.bold = True
                ab = self._para(align="justify" if f.justify else "left", space_after=6)
                self._runs(ab, p.abstract)
        if p.keywords:
            kp = self._para(align="left", space_after=10)
            lab = kp.add_run(f.keywords_label + ("—" if f.id == "ieee" else " " if f.keywords_label.endswith(":") else ": "))
            lab.bold = f.id != "ieee"
            lab.italic = f.id in ("ieee", "apa7")
            kr = kp.add_run(f.keywords_sep.join(p.keywords) + ("." if f.id == "ieee" else ""))
            kr.bold = f.id == "ieee" or None
        if p.highlights and f.id == "elsevier":
            self._para().add_run("Highlights").bold = True
            for h in p.highlights:
                self._para(style="List Bullet").add_run(h)
        if f.title_page:  # APA: the text starts on a new page under the title
            self.doc.paragraphs[-1].add_run().add_break(WD_BREAK.PAGE)
            tp = self._para(align="center")
            tp.add_run(p.title).bold = True

    # ------------------------------------------------------------------ body
    def _heading(self, text: str, level: int):
        f = self.f
        h = self.doc.add_heading(level=min(3, level))
        if level == 1:
            h.alignment = ALIGN[f.h1_align]
        run = h.add_run(text)
        if level == 1 and f.h1_case == "smallcaps":
            run.font.small_caps = True
        return h

    def _block(self, b: Block, first_in_section: bool):
        f = self.f
        if b.kind in ("para", "quote"):
            p = self._para(indent=0 if first_in_section and f.id not in ("apa7",) else f.para_indent_in)
            if b.kind == "quote":
                p.paragraph_format.left_indent = Inches(0.4)
            self._runs(p, b.runs)
            return
        if b.kind == "code":
            p = self._para(align="left")
            r = p.add_run("".join(x.t for x in b.runs))
            r.font.name, r.font.size = "Consolas", Pt(f.body_pt - 1.5)
            return
        if b.kind == "list":
            for it in b.items:
                p = self._para(style="List Number" if b.ordered else "List Bullet")
                self._runs(p, it)
            return
        if b.kind == "equation":
            self._equation(b)
            return
        if b.kind in ("figure", "table"):
            self._float(b)

    def _equation(self, b: Block):
        n = self.lay.eq_no.get(id(b), 0)
        # equation centred with the number at the right margin (tab stops)
        p = self._para(align="left")
        pf = p.paragraph_format
        pf.space_before, pf.space_after = Pt(4), Pt(4)
        pf.tab_stops.add_tab_stop(Inches(self.col_w / 2), alignment=1)  # centre
        pf.tab_stops.add_tab_stop(Inches(self.col_w - 0.02), alignment=2)  # right
        p.add_run("\t")
        if b.omml:
            el = parse_xml(b.omml)
            maths = el.findall(qn("m:oMath")) if el.tag == qn("m:oMathPara") else [el]
            for m in maths:
                p._p.append(copy.deepcopy(m))
        elif b.latex and (res := render_math(b.latex, self.math_dir, size_pt=self.f.body_pt)):
            path, w, h = res
            p.add_run().add_picture(str(path), width=Pt(min(w, self.col_w * 72 * 0.85)))
        elif b.image and (self.dir / b.image).exists():
            from PIL import Image

            with Image.open(self.dir / b.image) as im:
                w_pt = im.width / 4  # equations are cropped at 4x (288 dpi)
            p.add_run().add_picture(str(self.dir / b.image), width=Pt(min(w_pt, self.col_w * 72 * 0.85)))
        else:
            p.add_run(as_text(b.latex) if b.latex else "".join(r.t for r in b.runs)).italic = True
            self.math_as_text += 1
        p.add_run(f"\t({n})")

    def _float(self, b: Block):
        f = self.f
        label = caption_label(f, b, self.lay)

        def caption():
            if f.caption_newline:  # APA: bold label, then the title in italics on the next line
                c = self._para(style="Caption", align="left")
                c.add_run(label).bold = True
                c.add_run().add_break()
                self._runs(c, b.caption, size=f.caption_pt, italic=True)
                return c
            c = self._para(style="Caption", align="center" if f.id in ("ieee", "springer_lncs") or b.kind == "figure" else "left")
            c.paragraph_format.space_after = Pt(6 if b.kind == "figure" else 3)
            c.paragraph_format.space_before = Pt(3)
            lr = c.add_run(label + (" " if f.fig_sep == " " else ""))
            lr.bold = f.id not in ("ieee",)
            if f.table_roman and b.kind == "table":
                lr.font.small_caps = True
                c.add_run("\n")
            self._runs(c, b.caption, size=f.caption_pt)
            if f.table_roman and b.kind == "table":
                for r in c.runs[2:]:
                    r.font.small_caps = True
            return c

        above = (b.kind == "table" and f.table_caption_above) or (b.kind == "figure" and f.fig_caption_above)
        if above:
            caption()
        if b.kind == "table" and b.rows:
            self._table(b.rows)
        elif b.image and (self.dir / b.image).exists():
            p = self._para(align="center")
            p.paragraph_format.keep_with_next = True
            from PIL import Image

            with Image.open(self.dir / b.image) as im:
                ar = im.height / im.width
            w = self.col_w * 0.98
            if w * ar > 6.5:
                w = 6.5 / ar
            p.add_run().add_picture(str(self.dir / b.image), width=Inches(w))
        else:
            self._para(align="center").add_run(f"[{label}: file not available]").italic = True
        if not above:
            caption()

    def _table(self, rows: list[list[str]]):
        ncol = max(len(r) for r in rows)
        t = self.doc.add_table(rows=len(rows), cols=ncol)
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        for i, r in enumerate(rows):
            for j in range(ncol):
                cell = t.cell(i, j)
                cell.text = ""
                para = cell.paragraphs[0]
                run = para.add_run(r[j] if j < len(r) else "")
                run.font.size = Pt(self.f.caption_pt)
                run.bold = i == 0 or None
        # booktabs-like rules: top, below header, bottom
        tbl_pr = t._tbl.tblPr
        borders = OxmlElement("w:tblBorders")
        for edge, sz in (("top", 12), ("bottom", 12)):
            e = OxmlElement(f"w:{edge}")
            e.set(qn("w:val"), "single")
            e.set(qn("w:sz"), str(sz))
            e.set(qn("w:color"), "000000")
            borders.append(e)
        tbl_pr.append(borders)
        for cell in t.rows[0].cells:
            tcpr = cell._tc.get_or_add_tcPr()
            b = OxmlElement("w:tcBorders")
            e = OxmlElement("w:bottom")
            e.set(qn("w:val"), "single")
            e.set(qn("w:sz"), "6")
            b.append(e)
            tcpr.append(b)
        self._para(space_after=4)

    # ------------------------------------------------------------------ references
    def _references(self):
        f = self.f
        h = self._heading(f.references_title.upper() if f.h1_case in ("upper", "smallcaps") else f.references_title, 1)
        if f.id == "apa7":
            h.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for k in self.r.order:
            ref = self.p.ref(k)
            p = self._para(align="left")
            pf = p.paragraph_format
            if self.r.author_year:
                pf.left_indent, pf.first_line_indent = Inches(0.5 if f.id == "apa7" else 0.25), Inches(-(0.5 if f.id == "apa7" else 0.25))
            else:
                pf.left_indent, pf.first_line_indent = Inches(0.33), Inches(-0.33)
                pf.tab_stops.add_tab_stop(Inches(0.33))
                lab = p.add_run(self.r.labels.get(k, "") + "\t")
                lab.font.size = Pt(f.ref_pt)
            pf.space_after = Pt(2 if f.line_spacing < 1.5 else 0)
            self._runs(p, self.r.entries.get(k) or [Inline(t=ref.raw if ref else k)], size=f.ref_pt)

    # ------------------------------------------------------------------ assemble
    def write(self, out: Path) -> Path:
        f = self.f
        self._setup()
        self._front()
        if f.columns == 2:
            sec = self.doc.add_section(WD_SECTION.CONTINUOUS)
            cols = sec._sectPr.xpath("./w:cols")
            c = cols[0] if cols else OxmlElement("w:cols")
            c.set(qn("w:num"), "2")
            c.set(qn("w:space"), str(int(f.column_gap_in * 1440)))
            if not cols:
                sec._sectPr.append(c)
        for n in self.lay.body:
            if n.heading:
                self._heading(n.heading, n.section.level)
            for i, b in enumerate(n.section.blocks):
                self._block(b, i == 0)
        if self.lay.acknowledgements:
            self._heading(ack_title(f), 1 if f.id != "springer_lncs" else 2)
            for s in self.lay.acknowledgements:
                for i, b in enumerate(s.blocks):
                    self._block(b, i == 0)
        stmts = {k: s for k, s in self.lay.statements.items()}
        if stmts or self.lay.missing:
            if f.id == "springer_nature":
                self._heading("Declarations", 1)
                for k in ("funding", "competing", "ethics", "data", "contributions", "declarations"):
                    if k in stmts:
                        p = self._para(style="List Bullet")
                        if k != "declarations":
                            p.add_run(statement_title(f, k) + ": ").bold = True
                        for b in stmts[k].blocks:
                            self._runs(p, b.runs)
                    elif k in self.lay.missing:
                        p = self._para(style="List Bullet")
                        p.add_run(statement_title(f, k) + ": ").bold = True
                        r = p.add_run(PLACEHOLDER)
                        r.font.highlight_color = 7  # yellow
            else:
                for k, s in stmts.items():
                    self._heading(statement_title(f, k) if k != "declarations" else "Declarations", 1)
                    for i, b in enumerate(s.blocks):
                        self._block(b, i == 0)
                for k in self.lay.missing:
                    self._heading(statement_title(f, k), 1)
                    r = self._para().add_run(PLACEHOLDER)
                    r.font.highlight_color = 7
        if self.lay.appendix:
            for n in self.lay.appendix:
                if n.heading:
                    self._heading(("Appendix " + n.number + " " + n.section.title) if n.section.level == 1 and not n.section.title.lower().startswith("appendix") else n.section.title, n.section.level)
                for i, b in enumerate(n.section.blocks):
                    self._block(b, i == 0)
        if self.p.references:
            self._references()
        out.parent.mkdir(parents=True, exist_ok=True)
        self.doc.save(out)
        return out


def write(paper: Paper, fmt: Format, rend: Rendered, paper_dir: Path, out: Path, style: str = "") -> tuple[Path, int]:
    w = WordWriter(paper, fmt, rend, paper_dir, style)
    w.write(out)
    return out, w.math_as_text
