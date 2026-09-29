"""Video editor API: the editable timeline, script import and preview/final rendering."""
from __future__ import annotations

import json
import tempfile
from pathlib import Path
from typing import Any, Literal

from fastapi import APIRouter, File, HTTPException, UploadFile
from pydantic import BaseModel, ValidationError

from backend.media.library import MediaLibrary
from backend.pipeline.jobs import runner
from backend.pipeline.orchestrator import STAGES_BY_KIND, _theme
from backend.pipeline.outputs import project_work_dir
from backend.pipeline.video import ensure_timeline
from backend.rendering.video.timeline import Timeline, save_timeline
from backend.storage.database import get_db

router = APIRouter(prefix="/api/projects", tags=["video"])


class RenderRequest(BaseModel):
    quality: Literal["preview", "final"] = "final"


def _project_and_plan(project_id: str):
    db = get_db()
    p = db.get_project(project_id)
    if not p:
        raise HTTPException(404, "Project not found")
    plan = db.latest_plan(project_id)
    if not plan:
        raise HTTPException(400, "Generate the presentation first.")
    return db, p, plan


def _enrich(tl: Timeline, project: dict) -> dict[str, Any]:
    lib = MediaLibrary()
    data = json.loads(tl.model_dump_json())
    for c in data["clips"]:
        a = lib.get(c["asset_id"]) if c.get("asset_id") else None
        c["asset"] = {"id": a.id, "kind": a.kind, "filename": a.filename, "url": f"/api/media/{a.id}/file",
                      "thumb_url": f"/api/media/{a.id}/thumb" if a.thumb else None, "duration": a.meta.get("duration"),
                      "width": a.width, "height": a.height, "has_audio": a.meta.get("has_audio", False)} if a else None
        if c.get("narration_asset_id"):
            n = lib.get(c["narration_asset_id"])
            c["narration_asset"] = {"id": n.id, "filename": n.filename, "duration": n.meta.get("duration")} if n else None
        if c["type"] == "slide":
            c["thumb_url"] = f"/api/projects/{project['id']}/files/slides/slide_{c['slide_number']:02d}.png"
    for key in ("music", "narration"):
        aid = data[key].get("asset_id")
        a = lib.get(aid) if aid else None
        data[key]["asset"] = {"id": a.id, "filename": a.filename, "duration": a.meta.get("duration")} if a else None
    work = project_work_dir(project)
    rendered = {}
    for q in ("preview", "final"):
        f = work / f"rendered_{q}.json"
        if f.exists():
            rendered[q] = json.loads(f.read_text(encoding="utf-8"))
    data["rendered"] = rendered
    return data


@router.get("/{project_id}/timeline")
def get_timeline(project_id: str):
    db, p, plan = _project_and_plan(project_id)
    return _enrich(ensure_timeline(db, p, plan, _theme(p)), p)


@router.put("/{project_id}/timeline")
def put_timeline(project_id: str, body: dict[str, Any]):
    db, p, plan = _project_and_plan(project_id)
    body = {k: v for k, v in body.items() if k != "rendered"}
    for c in body.get("clips", []):
        c.pop("asset", None), c.pop("thumb_url", None), c.pop("narration_asset", None)
    for key in ("music", "narration"):
        if isinstance(body.get(key), dict):
            body[key].pop("asset", None)
    try:
        tl = Timeline.model_validate(body)
    except ValidationError as exc:
        raise HTTPException(422, "; ".join(f"{'.'.join(map(str, e['loc']))}: {e['msg']}" for e in exc.errors()[:4])) from exc
    lib = MediaLibrary()
    for c in tl.clips:
        if c.type in ("image", "video"):
            a = lib.get(c.asset_id or "")
            if not a or a.kind != c.type:
                raise HTTPException(422, f"Clip '{c.label or c.type}' refers to a missing {c.type}")
            if c.type == "video":
                c.source_duration = float(a.meta.get("duration") or c.source_duration)
                if c.trim_end is not None and c.trim_end <= c.trim_start:
                    raise HTTPException(422, f"Clip '{c.label}': the trim end must be after the trim start")
        if c.type == "slide" and not any(s.slide_number == c.slide_number for s in plan.slides):
            raise HTTPException(422, f"Slide {c.slide_number} does not exist")
    for aid in (tl.music.asset_id, tl.narration.asset_id):
        if aid and not lib.get(aid):
            raise HTTPException(422, "The selected audio file is missing")
    save_timeline(project_work_dir(p), tl)
    return _enrich(tl, p)


@router.post("/{project_id}/timeline/reset")
def reset_timeline(project_id: str):
    db, p, plan = _project_and_plan(project_id)
    return _enrich(ensure_timeline(db, p, plan, _theme(p), rebuild=True), p)


@router.post("/{project_id}/video/render")
def render(project_id: str, body: RenderRequest):
    db, p, _ = _project_and_plan(project_id)
    if p.get("last_job_id"):
        last = db.get_job(p["last_job_id"])
        if last and last.status in ("queued", "running"):
            raise HTTPException(409, "A job for this project is already running.")
    if "mp4" not in (p["output_formats"] or []):
        db.update_project(project_id, output_formats=list(p["output_formats"] or []) + ["mp4"])
    job = runner.submit(project_id, "video", STAGES_BY_KIND["video"], {"quality": body.quality})
    return job.model_dump()


async def _script_text(file: UploadFile) -> str:
    name = (file.filename or "").lower()
    data = await file.read()
    if len(data) > 5 * 1024 * 1024:
        raise HTTPException(400, "The script is too large (5 MB max).")
    if name.endswith(".docx"):
        import docx

        with tempfile.NamedTemporaryFile(suffix=".docx", delete=False) as t:
            t.write(data)
        try:
            text = "\n".join(p.text for p in docx.Document(t.name).paragraphs)
        finally:
            Path(t.name).unlink(missing_ok=True)
    elif name.endswith((".txt", ".md", ".markdown", ".srt", ".vtt")):
        text = data.decode("utf-8", errors="ignore")
        if name.endswith((".srt", ".vtt")):  # keep spoken lines only
            text = "\n".join(l for l in text.splitlines() if l.strip() and "-->" not in l and not l.strip().isdigit()
                             and l.strip() != "WEBVTT")
    else:
        raise HTTPException(400, "Upload the script as .txt, .md, .docx, .srt or .vtt")
    return text.strip()[:100_000]


@router.post("/{project_id}/script")
async def import_script(project_id: str, file: UploadFile = File(...)):
    """Read a narration script from .txt / .md / .docx / .srt / .vtt (returned for the editor; nothing is sent anywhere)."""
    if not get_db().get_project(project_id):
        raise HTTPException(404, "Project not found")
    return {"script": await _script_text(file)}


script_router = APIRouter(prefix="/api/scripts", tags=["video"])


@script_router.post("/parse")
async def parse_script_file(file: UploadFile = File(...)):
    """Same as above, before a project exists (New Project wizard)."""
    from backend.planning.script import parse_script

    text = await _script_text(file)
    marked, plain = parse_script(text)
    return {"script": text, "sections": len(marked) or len(plain)}
