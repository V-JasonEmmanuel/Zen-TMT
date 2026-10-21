"""Build the documentation slide decks (editable PowerPoint, Zensar colours).

    python scripts/make_docs_decks.py   ->  docs/Module_Architecture.pptx, docs/Pipeline.pptx
"""
from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
DOCS = ROOT / "docs"
LOGO_W = ROOT / "brands/zensar/assets/logo/zensar_logo_white.png"
LOGO_D = ROOT / "brands/zensar/assets/logo/zensar_logo_dark.png"
INDIGO, CORAL, BLUE, TEAL = RGBColor(0x10, 0x00, 0x5D), RGBColor(0xF0, 0x4E, 0x45), RGBColor(0x3A, 0x57, 0xA7), RGBColor(0x00, 0xB0, 0xC0)
LIGHT, GREY, INK, WHITE = RGBColor(0xF0, 0xF0, 0xEE), RGBColor(0x6D, 0x6E, 0x71), RGBColor(0x1A, 0x1A, 0x1A), RGBColor(0xFF, 0xFF, 0xFF)
SOFT = [RGBColor(0xD3, 0xE1, 0xF4), RGBColor(0xCF, 0xEB, 0xE8), RGBColor(0xF1, 0xCF, 0xD5), RGBColor(0xED, 0xDB, 0xAA)]
FONT = "Arial"
W, H = Inches(13.333), Inches(7.5)


class Deck:
    def __init__(self, title: str):
        self.p = Presentation()
        self.p.slide_width, self.p.slide_height = W, H
        self.blank = self.p.slide_layouts[6]
        self.n = 0
        self.title = title

    # ---------------- primitives
    def _bg(self, s, color):
        r = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, W, H)
        r.fill.solid(); r.fill.fore_color.rgb = color; r.line.fill.background()
        return r

    def text(self, s, x, y, w, h, txt, size=18, color=INK, bold=False, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP):
        tb = s.shapes.add_textbox(x, y, w, h)
        tf = tb.text_frame
        tf.word_wrap = True
        tf.vertical_anchor = anchor
        tf.margin_left = tf.margin_right = Inches(0.05)
        lines = txt if isinstance(txt, list) else [txt]
        for i, ln in enumerate(lines):
            para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            para.alignment = align
            run = para.add_run()
            run.text = ln
            run.font.size, run.font.bold, run.font.name = Pt(size), bold, FONT
            run.font.color.rgb = color
        return tb

    def box(self, s, x, y, w, h, title, body="", fill=SOFT[0], color=INK, title_size=15, body_size=11.5, shape=MSO_SHAPE.ROUNDED_RECTANGLE):
        r = s.shapes.add_shape(shape, x, y, w, h)
        r.fill.solid(); r.fill.fore_color.rgb = fill; r.line.fill.background()
        if shape == MSO_SHAPE.ROUNDED_RECTANGLE:
            r.adjustments[0] = 0.12
        tf = r.text_frame
        tf.word_wrap = True
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        tf.margin_left = tf.margin_right = Inches(0.12)
        p0 = tf.paragraphs[0]
        p0.alignment = PP_ALIGN.CENTER
        run = p0.add_run(); run.text = title
        run.font.size, run.font.bold, run.font.name = Pt(title_size), True, FONT
        run.font.color.rgb = color
        if body:
            for ln in body.split("\n"):
                p = tf.add_paragraph(); p.alignment = PP_ALIGN.CENTER
                r2 = p.add_run(); r2.text = ln
                r2.font.size, r2.font.name = Pt(body_size), FONT
                r2.font.color.rgb = color
        return r

    def arrow(self, s, x1, y1, x2, y2, color=GREY):
        c = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, x1, y1, x2, y2)
        c.line.color.rgb = color
        c.line.width = Pt(2)
        ln = c.line._get_or_add_ln()
        from lxml import etree
        tail = etree.SubElement(ln, "{http://schemas.openxmlformats.org/drawingml/2006/main}tailEnd")
        tail.set("type", "triangle")
        return c

    def footer(self, s, dark=False):
        self.n += 1
        self.text(s, Inches(0.6), Inches(7.0), Inches(9), Inches(0.35), f"Zensar Content Studio  |  {self.title}", 10, WHITE if dark else GREY)
        self.text(s, Inches(11.9), Inches(7.0), Inches(0.9), Inches(0.35), str(self.n), 10, WHITE if dark else GREY, align=PP_ALIGN.RIGHT)

    # ---------------- slide types
    def cover(self, title, subtitle):
        s = self.p.slides.add_slide(self.blank)
        self._bg(s, INDIGO)
        for i, (x, y, d, c) in enumerate([(9.6, -1.2, 5.2, BLUE), (11.2, 4.4, 3.4, TEAL), (8.4, 5.6, 2.4, CORAL)]):
            o = s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(y), Inches(d), Inches(d))
            o.fill.solid(); o.fill.fore_color.rgb = c; o.line.fill.background()
        s.shapes.add_picture(str(LOGO_W), Inches(0.8), Inches(0.7), height=Inches(0.55))
        self.text(s, Inches(0.8), Inches(2.5), Inches(8.6), Inches(1.8), title, 44, WHITE, True)
        bar = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.85), Inches(4.35), Inches(1.1), Inches(0.09))
        bar.fill.solid(); bar.fill.fore_color.rgb = CORAL; bar.line.fill.background()
        self.text(s, Inches(0.8), Inches(4.65), Inches(8.4), Inches(1.4), subtitle, 20, WHITE)
        self.n += 1
        return s

    def section(self, title, kicker=""):
        s = self.p.slides.add_slide(self.blank)
        self._bg(s, BLUE)
        if kicker:
            self.text(s, Inches(0.8), Inches(2.6), Inches(11), Inches(0.5), kicker.upper(), 14, WHITE, True)
        self.text(s, Inches(0.8), Inches(3.1), Inches(11.5), Inches(1.4), title, 38, WHITE, True)
        self.footer(s, dark=True)
        return s

    def slide(self, title, subtitle=""):
        s = self.p.slides.add_slide(self.blank)
        self._bg(s, WHITE)
        self.text(s, Inches(0.6), Inches(0.45), Inches(11.4), Inches(0.8), title, 28, INDIGO, True)
        bar = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.65), Inches(1.22), Inches(0.75), Inches(0.07))
        bar.fill.solid(); bar.fill.fore_color.rgb = CORAL; bar.line.fill.background()
        if subtitle:
            self.text(s, Inches(0.6), Inches(1.38), Inches(12), Inches(0.5), subtitle, 14, GREY)
        logo = s.shapes.add_picture(str(LOGO_D), 0, Inches(0.5), height=Inches(0.3))
        logo.left = W - Inches(0.6) - logo.width
        self.footer(s)
        return s

    def bullets(self, title, items, subtitle=""):
        s = self.slide(title, subtitle)
        tb = s.shapes.add_textbox(Inches(0.75), Inches(2.0), Inches(11.8), Inches(4.8))
        tf = tb.text_frame; tf.word_wrap = True
        for i, it in enumerate(items):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            head, _, rest = it.partition(" — ")
            r1 = p.add_run(); r1.text = "●  "; r1.font.color.rgb = CORAL; r1.font.size = Pt(14); r1.font.name = FONT
            r2 = p.add_run(); r2.text = head; r2.font.bold = bool(rest); r2.font.size = Pt(18); r2.font.color.rgb = INK; r2.font.name = FONT
            if rest:
                r3 = p.add_run(); r3.text = " — " + rest; r3.font.size = Pt(18); r3.font.color.rgb = INK; r3.font.name = FONT
            p.space_after = Pt(10)
        return s

    def table(self, title, header, rows, widths, subtitle="", size=12):
        s = self.slide(title, subtitle)
        top = Inches(2.0 if subtitle else 1.7)
        shape = s.shapes.add_table(len(rows) + 1, len(header), Inches(0.6), top, Inches(12.1), Inches(0.4) * (len(rows) + 1))
        t = shape.table
        for j, wd in enumerate(widths):
            t.columns[j].width = Inches(wd)
        for j, hcell in enumerate(header):
            c = t.cell(0, j); c.text = hcell
            c.fill.solid(); c.fill.fore_color.rgb = INDIGO
            for p in c.text_frame.paragraphs:
                for r in p.runs:
                    r.font.bold, r.font.size, r.font.color.rgb, r.font.name = True, Pt(size + 1), WHITE, FONT
        for i, row in enumerate(rows, start=1):
            for j, v in enumerate(row):
                c = t.cell(i, j); c.text = v
                c.fill.solid(); c.fill.fore_color.rgb = LIGHT if i % 2 else WHITE
                for p in c.text_frame.paragraphs:
                    for r in p.runs:
                        r.font.size, r.font.color.rgb, r.font.name = Pt(size), INK, FONT
                        r.font.bold = j == 0
        return s

    def flow(self, title, steps, subtitle="", note=""):
        """Steps left-to-right (wrapping to a second row); each (title, body)."""
        s = self.slide(title, subtitle)
        per_row = 5 if len(steps) > 5 else len(steps)
        rows = (len(steps) + per_row - 1) // per_row
        bw, gap = Inches(2.15), Inches(0.38)
        bh = Inches(2.9 if rows == 1 else 1.75)
        x0 = (W - (bw * per_row + gap * (per_row - 1))) // 2
        y0 = Inches(2.45 if rows == 1 else 1.95)
        for i, (t, b) in enumerate(steps):
            r, c = divmod(i, per_row)
            if r % 2:
                c = per_row - 1 - c  # snake layout so arrows stay short
            x = x0 + (bw + gap) * c
            y = y0 + (bh + Inches(0.55)) * r
            num = s.shapes.add_shape(MSO_SHAPE.OVAL, x - Inches(0.12), y - Inches(0.12), Inches(0.42), Inches(0.42))
            num.fill.solid(); num.fill.fore_color.rgb = CORAL; num.line.fill.background()
            tf = num.text_frame; tf.margin_left = tf.margin_right = 0
            pr = tf.paragraphs[0]; pr.alignment = PP_ALIGN.CENTER
            rr = pr.add_run(); rr.text = str(i + 1); rr.font.size = Pt(12); rr.font.bold = True; rr.font.color.rgb = WHITE; rr.font.name = FONT
            self.box(s, x, y, bw, bh, t, b, fill=SOFT[i % 4], title_size=16 if rows == 1 else 14, body_size=13 if rows == 1 else 11)
            s.shapes._spTree.remove(num._element); s.shapes._spTree.append(num._element)
            if i + 1 < len(steps):
                r2, c2 = divmod(i + 1, per_row)
                if r2 % 2:
                    c2 = per_row - 1 - c2
                if r2 == r:
                    xa = x + bw if c2 > c else x
                    xb = x0 + (bw + gap) * c2 + (0 if c2 > c else bw)
                    self.arrow(s, xa, y + bh // 2, xb, y + bh // 2)
                else:
                    self.arrow(s, x + bw // 2, y + bh, x + bw // 2, y0 + (bh + Inches(0.55)) * r2)
        if note:
            self.text(s, Inches(0.6), Inches(6.35), Inches(12.1), Inches(0.6), note, 13, GREY)
        return s

    def save(self, path: Path):
        self.p.save(str(path))
        print("wrote", path)


# ====================================================================== Module Architecture deck
def architecture():
    d = Deck("Module Architecture")
    d.cover("Module Architecture", "How Zensar Content Studio is built:\nthe layers, the modules and what each one does")

    d.bullets("What the app does", [
        "Projects — documents, demo videos and code repositories become an editable PowerPoint, brand visuals and a narrated video",
        "Research Papers — a paper is re-formatted for Springer Nature, LNCS, Elsevier, IEEE, ACM or APA (LaTeX, Word, PDF, BibTeX)",
        "Brand Documents — any document becomes a branded white paper designed like a reference PDF, with images created for its content",
        "Private — everything runs on this computer; documents never leave it",
        "Faithful — every statement is traced to its source; numbers are checked",
    ], "Three tools, one local engine")

    # layered architecture diagram
    s = d.slide("Architecture overview", "Browser interface → local server → engine modules → local AI and storage")
    d.box(s, Inches(0.6), Inches(1.95), Inches(12.1), Inches(0.8), "Browser — React interface",
          "Dashboard · Projects · Research Papers · Brand Documents · Templates · Media Library · Brand Manager · Settings", fill=INDIGO, color=WHITE)
    d.arrow(s, Inches(6.65), Inches(2.75), Inches(6.65), Inches(3.0))
    d.box(s, Inches(0.6), Inches(3.0), Inches(12.1), Inches(0.7), "FastAPI server — api/ (REST endpoints)  +  pipeline/ (background jobs, orchestrator)",
          fill=BLUE, color=WHITE, title_size=14)
    d.arrow(s, Inches(6.65), Inches(3.7), Inches(6.65), Inches(3.95))
    mods = [("ingestion\nextraction", "read files"), ("intelligence", "search · verify"), ("llm\nplanning", "write · plan"),
            ("rendering", "slides · images · video"), ("papers", "publisher formats"), ("branddocs", "branded PDFs"),
            ("branding · media", "brand rules · assets"), ("storage · utils", "SQLite · privacy guard")]
    for i, (t, b) in enumerate(mods):
        x = Inches(0.6) + (Inches(1.47) + Inches(0.06)) * i
        d.box(s, x, Inches(3.95), Inches(1.47), Inches(1.25), t, b, fill=SOFT[i % 4], title_size=11.5, body_size=9.5)
    d.arrow(s, Inches(3.0), Inches(5.2), Inches(3.0), Inches(5.55))
    d.arrow(s, Inches(6.65), Inches(5.2), Inches(6.65), Inches(5.55))
    d.arrow(s, Inches(10.3), Inches(5.2), Inches(10.3), Inches(5.55))
    d.box(s, Inches(0.6), Inches(5.55), Inches(3.9), Inches(1.05), "Ollama (localhost)", "Qwen 2.5 writing AI · qwen2.5vl vision", fill=CORAL, color=WHITE)
    d.box(s, Inches(4.7), Inches(5.55), Inches(3.9), Inches(1.05), "Image worker (own process)", "Stable Diffusion · DreamShaper 8 LCM", fill=CORAL, color=WHITE)
    d.box(s, Inches(8.8), Inches(5.55), Inches(3.9), Inches(1.05), "Local files + SQLite", "data/ · outputs/ · projects/ · models/", fill=CORAL, color=WHITE)

    d.table("Interface modules (frontend/src)", ["Module", "What it does"], [
        ["pages/", "One screen each: Dashboard, NewProject, Projects, ProjectDetail, ResearchPapers, PaperDetail, BrandDocuments, BrandDocDetail, Templates, MediaLibrary, Brands, BrandEditor, Settings, Welcome"],
        ["components/", "Shared building blocks: buttons/cards/tabs (ui), FileDrop, JobProgress, Presenter, SlideEditor, VideoEditor, MediaPicker, TemplateGallery, OutputPanels"],
        ["layouts/AppLayout", "The frame: left menu, page header, system status"],
        ["services/api.ts", "Every call from the screens to the server"],
        ["App.tsx / main.tsx", "Which page shows for which address; starting the app"],
        ["dist/", "The ready-built interface, served by the server (no Node.js needed)"],
    ], [2.6, 9.5], "React 18 · TypeScript · Tailwind CSS · Vite")

    d.table("Server and API (backend/main.py, backend/api)", ["Module", "Responsibility"], [
        ["main.py", "Starts FastAPI/Uvicorn, registers every API, serves the interface, installs the privacy guard"],
        ["documents · projects · generation", "Upload sources, manage projects, start generation jobs and report progress"],
        ["papers · branddocs", "Research Papers and Brand Documents: convert, review, re-export, download"],
        ["brands · templates · media", "Brand profiles, template analysis, Media Library"],
        ["video", "Video editor: timeline, script import, rendering"],
        ["settings", "Settings and system status (AI, models, voices, FFmpeg, storage)"],
        ["pipeline/jobs · orchestrator", "Runs jobs one at a time in the background, stage by stage, with real progress"],
    ], [3.6, 8.5], "Every screen talks to one API module; long work runs as a background job")

    d.flow("Reading documents — ingestion & extraction", [
        ("Adapters", "pdf · docx · pptx · html\nmarkdown · txt · video · repo"),
        ("PDF reading", "PyMuPDF: columns, headers,\nhyphens, tables, figures"),
        ("OCR (optional)", "RapidOCR for scanned pages"),
        ("Structure", "sections, section types,\ntitle"),
        ("Result", "typed blocks with page\nnumbers (provenance)"),
    ], "backend/ingestion · backend/extraction")

    d.flow("Understanding — intelligence", [
        ("Chunking", "small passages that keep\npage and section"),
        ("Embeddings", "MiniLM on ONNX Runtime:\nmeaning as numbers"),
        ("Retrieval", "FAISS vector search"),
        ("Relevance", "instruction + structure +\nsemantics + keywords"),
        ("Verification", "every claim checked;\nnumbers must match"),
    ], "backend/intelligence")

    d.flow("Writing and planning — llm & planning", [
        ("Local AI", "Ollama · Qwen 2.5\n(localhost only)"),
        ("Prompts", "versioned templates\nllm/prompts/*.txt"),
        ("Structured output", "JSON → validate (Pydantic)\n→ repair / retry"),
        ("Content contract", "instruction → precise\nspecification"),
        ("Planners", "content · slides · visuals\n· video · script"),
    ], "backend/llm · backend/planning")

    d.flow("Drawing — rendering", [
        ("Layout engine", "semantic layout + brand\ntokens → positions"),
        ("Scene graph", "one drawing description\nfor every output"),
        ("PowerPoint", "python-pptx: native,\neditable slides"),
        ("Images", "Pillow / Matplotlib PNG\n+ SVG vectors"),
        ("Video", "motion, narration (SAPI),\nmusic, FFmpeg → MP4"),
    ], "backend/rendering")

    d.table("Brand and media modules", ["Module", "What it does"], [
        ["branding/brand_profile · theme", "Brand values from official sources; missing values get labelled neutral fall-backs, never guesses"],
        ["branding/presets · template_parser", "Design presets from brand colours; reading PowerPoint templates"],
        ["branding/social · reference_analyzer", "Zensar 'Experience' design language; colour suggestions from brand images"],
        ["rendering/motifs · experience · grid", "Zensar shape language, cover/divider designs, modular grid"],
        ["media/library · treatment", "Media Library (stored once, with licence); crops, tints, borders"],
        ["media/online · query (optional)", "Licensed image search, off by default; only generic words are sent"],
        ["vision/keyframes · describe", "Scenes of demo videos; descriptions by a local vision model"],
    ], [3.9, 8.2])

    d.flow("Research Papers module (backend/papers)", [
        ("Readers", "PDF · Word · LaTeX ·\nMarkdown / text"),
        ("Structure", "title, authors, abstract,\nsections, statements"),
        ("References + citations", "rule-based parsing;\nlinks [1] / (Smith, 2020)"),
        ("Formats + CSL", "publisher rules;\ncitations by citeproc-py"),
        ("Writers + checks", "LaTeX · Word · PDF ·\nBibTeX · report"),
    ], "service.py runs the job; model.py is the shared paper description")

    d.flow("Brand Documents module (backend/branddocs)", [
        ("reference.py", "learns the template from\na reference PDF"),
        ("content.py", "reads the document into\nsections"),
        ("fit.py", "condenses / expands to the\nreference length (local AI)"),
        ("imagegen + worker", "scene descriptions →\nStable Diffusion → grading"),
        ("compose.py", "lays out pages like the\nreference (ReportLab)"),
    ], "service.py runs the job; model.py describes a template; fonts.py picks Graphik or the fall-back")

    d.table("Storage, utilities and privacy", ["Module", "What it does"], [
        ["storage/database · models", "Local SQLite database: projects, jobs, settings"],
        ["utils/config", "Settings from .env: folders, port, model"],
        ["utils/network_guard", "Blocks every outgoing connection except this computer"],
        ["utils/logging", "Logs counts and timings only, never document text"],
        ["utils/files · progress", "Safe file paths; progress messages for long steps"],
        ["services/documents · github", "Read an upload once and cache it; GitHub import on request only"],
    ], [3.6, 8.5])

    d.table("AI models and key dependencies", ["Component", "Technology", "Notes"], [
        ["Writing AI", "Qwen 2.5 7B / 3B via Ollama", "Apache-2.0; 7B chosen on GPUs with 8 GB+"],
        ["Image AI", "DreamShaper 8 LCM (SD 1.5) · Diffusers · PyTorch", "OpenRAIL-M; GPU (CUDA 12.4) or CPU"],
        ["Meaning search", "all-MiniLM-L6-v2 · ONNX Runtime · FAISS", "90 MB, CPU"],
        ["Vision (optional)", "qwen2.5vl:3b via Ollama", "describes demo-video scenes"],
        ["Documents", "PyMuPDF · python-docx · python-pptx · BeautifulSoup", "reading every format"],
        ["Outputs", "ReportLab · python-pptx · Pillow · Matplotlib · FFmpeg", "PDF, PowerPoint, images, video"],
        ["Papers", "citeproc-py + CSL · bibtexparser · pylatexenc · Jinja2", "official publisher styles"],
        ["Server / UI", "FastAPI · Uvicorn · Pydantic · SQLite / React · Tailwind · Vite", "local web app"],
    ], [2.3, 5.6, 4.2], size=11)

    d.flow("Deployment — start.bat", [
        ("Checks", "port, VC++ runtime,\nPython 3.10–3.12"),
        ("Packages", "pip online, or the\nGitHub offline bundle"),
        ("Models + UI", "embeddings; prebuilt\ninterface or npm build"),
        ("AI", "GPU → CUDA PyTorch +\nimage model; Ollama + Qwen"),
        ("Start", "system check, server,\nbrowser at :8000"),
    ], "scripts/bootstrap.ps1", "Behind a corporate proxy, packages and models come from github.com/V-JasonEmmanuel/Zen-TMT-offline over git, checked with SHA-256.")

    s = d.section("Every module is local, traceable and replaceable", "Summary")
    d.save(DOCS / "Module_Architecture.pptx")


# ====================================================================== Pipeline deck
def pipeline():
    d = Deck("Pipelines")
    d.cover("How the Pipelines Work", "From a document to slides, a publisher-ready paper\nor a branded white paper - step by step")

    d.flow("Starting the app — start.bat", [
        ("Port", "already running? another\nprogram? → next free port"),
        ("Python", "find or install 3.12;\nprivate .venv"),
        ("Packages", "pip, or the git offline\nbundle (SHA-256 checked)"),
        ("Models + UI", "embedding model;\ninterface"),
        ("Image AI", "nvidia-smi → CUDA or CPU\nPyTorch; DreamShaper"),
        ("Writing AI", "Ollama + Qwen 2.5\n(7B on 8 GB+ GPU)"),
        ("Check", "check_environment.py"),
        ("Run", "server on :8000;\nbrowser opens"),
    ], "scripts/bootstrap.ps1 - safe to run again: finished steps are skipped")

    d.flow("Step one for every tool: reading", [
        ("Upload", "PDF · Word · PowerPoint ·\ntext · video · repository"),
        ("Layout", "PyMuPDF: text pieces with\nposition, size, font"),
        ("Clean-up", "columns, headers/footers,\nhyphens, drop caps"),
        ("Elements", "headings, paragraphs, lists,\ntables, figures, formulas"),
        ("Structured content", "sections with page\nnumbers (provenance)"),
    ], "backend/ingestion · backend/papers/readers")

    d.section("Projects: documents → slides, visuals, video", "Pipeline 1")
    d.flow("Projects pipeline", [
        ("Chunk", "passages that keep page\nand section"),
        ("Embed + search", "MiniLM (ONNX) +\nFAISS"),
        ("Relevance", "only what the instruction\nasks for"),
        ("Contract", "instruction → precise\nspecification"),
        ("Plan", "Qwen 2.5 writes slides from\nthe passages (JSON)"),
        ("Verify", "claims vs sources;\nnumbers must match"),
        ("Visuals", "charts from real tables;\ndiagrams from structure"),
        ("Render", "PPTX · PDF · PNG/SVG ·\nMP4 with narration"),
    ], "backend/pipeline/orchestrator.py", "Without a local AI the plan uses extractive mode: real sentences from the document.")
    d.table("Projects — outputs", ["Output", "Built with", "Where"], [
        ["Editable PowerPoint", "python-pptx from the scene graph", "Presentation tab · outputs/<project>/presentation/"],
        ["PDF + slide previews", "Pillow raster backend", "Presentation tab"],
        ["Visuals (PNG / SVG)", "components + Matplotlib charts", "Images tab"],
        ["Narrated video (MP4)", "motion graphics, SAPI voices, music, FFmpeg", "Video editor tab"],
        ["Source mapping", "verification links", "Sources tab"],
    ], [3.0, 4.6, 4.5])

    d.section("Research Papers: paper → publisher format", "Pipeline 2")
    d.flow("Research Papers pipeline", [
        ("Read", "PDF / Word / LaTeX /\ntext readers"),
        ("Structure", "title, authors, abstract,\nsections, statements"),
        ("References", "split + parse by rules;\nAI only as checked fallback"),
        ("Citations", "[1], [2–5], (Smith, 2020)\nlinked to references"),
        ("Cross-refs", "Table 2, Fig. 3, Section 2.1\nrenumbered"),
        ("Format", "citeproc-py + official\nCSL styles"),
        ("Write", "LaTeX · Word · PDF ·\nBibTeX"),
        ("Report", "publisher checks +\ntext preserved"),
    ], "backend/papers/service.py")
    d.table("Supported publisher formats", ["Format", "LaTeX class", "Layout"], [
        ["Springer Nature journal", "sn-jnl", "single column, Declarations section"],
        ["Springer LNCS", "llncs", "proceedings text block"],
        ["Elsevier journal", "elsarticle", "preprint, highlights, CRediT"],
        ["IEEE", "IEEEtran", "two columns, Roman headings, TABLE I"],
        ["ACM conference", "acmart (sigconf)", "two columns, CCS concepts"],
        ["APA 7th edition", "apa7 + biblatex-apa", "title page, double spacing"],
    ], [3.5, 3.2, 5.4])

    d.section("Brand Documents: document → branded white paper", "Pipeline 3")
    d.flow("Learning a template from a reference PDF", [
        ("Text pieces", "type sizes: headings,\nbody, labels, footer"),
        ("Vector drawings", "cards, bands, rules, dashes,\nbullets, logo, pattern"),
        ("Rendered pixels", "true colours\n(blend modes included)"),
        ("Page roles", "cover · intro · text · image\n· conclusion · back"),
        ("Template", "geometry, styles, artwork,\nfooter, size profile"),
    ], "backend/branddocs/reference.py - done once per reference")
    d.flow("Creating a branded PDF", [
        ("Read", "sections, lists, tables,\nconclusion, authors"),
        ("Fit length", "Qwen condenses/expands to\nthe reference size"),
        ("Check numbers", "every number must be in\nthe source; trim to size"),
        ("Describe images", "AI writes a scene per\nimage from the content"),
        ("Generate", "Stable Diffusion in a\nseparate worker (GPU/CPU)"),
        ("Grade", "colours matched to the\nreference photos"),
        ("Lay out", "ReportLab: pages like the\nreference, vector artwork"),
        ("Check", "words arrived; report"),
    ], "backend/branddocs/service.py")
    d.table("Brand Documents — page designs", ["Page", "What goes on it"], [
        ["Cover", "generated image, logo (exact vectors), title, label (e.g. White Paper)"],
        ["Intro", "opening section in the coloured card; the next section below in two columns"],
        ["Text pages", "two columns, headings with the brand dash, rules, bullets, tables, formulas"],
        ["Image page", "photo fading into the page, one section per column"],
        ["Conclusion", "closing card or band (text above it when the reference does that)"],
        ["Back cover", "authors, brand pattern and logo, company text"],
    ], [2.4, 9.7])

    d.bullets("Trust built into every pipeline", [
        "Local only — a network guard blocks every outside connection; the AI runs on this computer",
        "Traceable — every statement keeps a link to its source page",
        "No invented numbers — numbers must appear in the source, or the text is replaced by the source's own sentences",
        "Checked results — every job ends with a report: words preserved, publisher checks, fonts, images",
        "Re-usable — papers are read once and exported to any format; images are cached",
    ])
    d.section("Upload a document - get finished, on-brand material", "Summary")
    d.save(DOCS / "Pipeline.pptx")


if __name__ == "__main__":
    DOCS.mkdir(exist_ok=True)
    architecture()
    pipeline()
