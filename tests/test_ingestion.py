"""Document processing: PDF / DOCX / TXT / Markdown / HTML / PPTX, malformed input, safety."""
from pathlib import Path

import pytest

from backend.ingestion import ExtractionError, get_adapter, supported_extensions
from backend.schemas import ContentType, SectionType
from backend.utils.files import UnsafePathError, safe_join, sanitize_filename


def test_pdf_extraction_structure_and_tables(extracted):
    doc, st, chunks = extracted
    assert doc.title.startswith("Adaptive Retrieval Pipelines")
    assert doc.page_count == 2
    assert doc.extraction_method == "native"
    headings = [b.text for b in doc.blocks if b.type == ContentType.heading]
    assert "4 Methodology" in headings and "5 Results" in headings
    tables = [b for b in doc.blocks if b.type == ContentType.table]
    assert len(tables) == 1 and tables[0].table[0][0] == "Configuration"
    assert tables[0].table_index == 1  # numbered from its "Table 1:" caption
    types = set(st.detected_types)
    assert {SectionType.abstract, SectionType.methodology, SectionType.results, SectionType.references} <= types


def test_chunks_keep_provenance(extracted):
    _, _, chunks = extracted
    assert chunks
    for c in chunks:
        assert c.document_id == "doc_test" and c.page >= 1 and c.section and c.text
        ref = c.ref()
        assert ref.page == c.page and ref.chunk_id == c.id
    table = next(c for c in chunks if c.content_type == ContentType.table)
    assert "Table 1" in table.source_reference and table.section.endswith("Results")


def test_docx_extraction(sample_docx, tmp_path):
    doc = get_adapter(sample_docx).extract(sample_docx, "d1", tmp_path)
    assert doc.title.startswith("Adaptive Retrieval")
    assert any(b.type == ContentType.table for b in doc.blocks)
    assert any(b.type == ContentType.heading and b.text == "Methodology" for b in doc.blocks)
    assert any(b.type == ContentType.list_item for b in doc.blocks)


def test_txt_markdown_html(tmp_path):
    (tmp_path / "a.txt").write_text("PROJECT OVERVIEW\n\nThe project reduced costs by 12%.\n\nNext Steps\n\n- Expand to two regions\n- Hire staff\n", encoding="utf-8")
    (tmp_path / "b.md").write_text("# Title\n\n## Results\n\nRevenue grew 8%.\n\n| A | B |\n|---|---|\n| x | 1 |\n\n```\ncode()\n```\n", encoding="utf-8")
    (tmp_path / "c.html").write_text("<html><head><title>T</title><script>alert(1)</script></head><body><h1>Intro</h1><p>Hello world text.</p>"
                                     "<table><tr><th>k</th><th>v</th></tr><tr><td>a</td><td>1</td></tr></table></body></html>", encoding="utf-8")
    t = get_adapter(tmp_path / "a.txt").extract(tmp_path / "a.txt", "t", tmp_path)
    assert t.blocks[0].type == ContentType.heading and any(b.type == ContentType.list_item for b in t.blocks)
    m = get_adapter(tmp_path / "b.md").extract(tmp_path / "b.md", "m", tmp_path)
    kinds = [b.type for b in m.blocks]
    assert ContentType.table in kinds and ContentType.code in kinds
    h = get_adapter(tmp_path / "c.html").extract(tmp_path / "c.html", "h", tmp_path)
    assert all("alert" not in b.text for b in h.blocks)  # scripts discarded, never executed
    assert any(b.type == ContentType.table for b in h.blocks)


def test_pptx_input(tmp_path):
    from pptx import Presentation
    from pptx.util import Inches

    prs = Presentation()
    s = prs.slides.add_slide(prs.slide_layouts[1])
    s.shapes.title.text = "Quarterly Review"
    s.placeholders[1].text_frame.text = "Revenue increased 14% year on year"
    path = tmp_path / "deck.pptx"
    prs.save(path)
    doc = get_adapter(path).extract(path, "p", tmp_path)
    assert doc.page_count == 1
    assert any(b.text == "Quarterly Review" for b in doc.blocks)
    assert any("14%" in b.text for b in doc.blocks)
    assert Inches  # imported for completeness


@pytest.mark.parametrize("name,data,msg", [
    ("broken.pdf", b"%PDF-1.4 this is not really a pdf", "could not be opened"),
    ("empty.txt", b"", "empty"),
    ("binary.txt", b"\x00\x01\x02\x00" * 100, "binary"),
    ("broken.docx", b"PK\x03\x04garbage", "could not be opened"),
])
def test_malformed_documents(tmp_path, name, data, msg):
    p = tmp_path / name
    p.write_bytes(data)
    with pytest.raises(ExtractionError) as e:
        get_adapter(p).extract(p, "x", tmp_path)
    assert msg in str(e.value).lower()


def test_unsupported_extension(tmp_path):
    with pytest.raises(ExtractionError):
        get_adapter(tmp_path / "file.exe")
    assert ".pdf" in supported_extensions() and ".exe" not in supported_extensions()


def test_filename_sanitising_and_traversal(tmp_path):
    assert sanitize_filename("../../etc/passwd") == "passwd"
    assert sanitize_filename("..\\..\\Windows\\win.ini") == "win.ini"
    assert sanitize_filename("CON.txt") == "_CON.txt"
    assert sanitize_filename("re port<>|.pdf") == "re port_.pdf"
    assert safe_join(tmp_path, "a", "b.txt").is_relative_to(tmp_path)
    with pytest.raises(UnsafePathError):
        safe_join(tmp_path, "..", "outside.txt")
    with pytest.raises(UnsafePathError):
        safe_join(tmp_path, str(Path("C:/Windows/win.ini")))
