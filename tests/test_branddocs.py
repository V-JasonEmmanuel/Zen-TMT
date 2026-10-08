"""Brand documents: learn a template from a reference PDF, convert text / Word documents into it."""
from __future__ import annotations

import re
from pathlib import Path

import pytest

LOREM = ("Teams often get stuck arguing about platforms long before they agree on the actual problem they are trying to solve. "
         "Value comes from understanding what matters, who is affected and why it matters now. ")


def _reference(path: Path) -> Path:
    """A small branded reference: image cover, card page, two-column text page, back cover."""
    import numpy as np
    from PIL import Image
    from reportlab.lib.colors import HexColor
    from reportlab.pdfgen import canvas

    W, H = 595.28, 841.89
    img = path.with_suffix(".png")
    y = np.linspace(0, 1, 400)[:, None, None]
    arr = (np.array([40, 30, 120]) * (1 - y) + np.array([5, 5, 20]) * y) * np.ones((1, 300, 1))
    arr[60:180, 100:200] = [230, 230, 255]
    Image.fromarray(arr.astype("uint8")).save(img)
    c = canvas.Canvas(str(path), pagesize=(W, H))
    # 1 cover
    c.drawImage(str(img), 0, 0, W, H)
    c.setFillColor(HexColor("#FFFFFF"))
    for i in range(6):  # logo lockup
        c.rect(440 + i * 20, H - 60, 14, 18, stroke=0, fill=1)
    c.setFont("Helvetica-Bold", 44)
    c.drawString(230, H - 520, "Branded title")
    c.drawString(230, H - 566, "for testing")
    c.setFont("Helvetica", 17)
    c.drawString(250, H - 600, "White Paper")
    c.setFillColor(HexColor("#EE4E46"))
    c.wedge(230, H - 604, 246, H - 588, 0, 90, stroke=0, fill=1)
    c.showPage()

    def body_lines(x, y0, n, color):
        c.setFillColor(HexColor(color))
        c.setFont("Helvetica", 10)
        words = (LOREM * 4).split()
        k = 0
        for i in range(n):
            c.drawString(x, H - (y0 + i * 15), " ".join(words[k:k + 7]))
            k += 7

    def footer(n, color):
        c.setStrokeColor(HexColor(color))
        c.setLineWidth(0.25)
        c.line(36, H - 806, 559, H - 806)
        c.setFillColor(HexColor(color))
        c.setFont("Helvetica", 7.8)
        c.drawRightString(560, H - 822, f"© Test Company, 2026  |  Page {n}")

    # 2 intro: dark page, coral card
    c.setFillColor(HexColor("#201B5A"))
    c.rect(0, 0, W, H, stroke=0, fill=1)
    c.setFillColor(HexColor("#ED4E46"))
    c.roundRect(36, H - 458, 523, 389, 30, stroke=0, fill=1)
    c.setFillColor(HexColor("#FFFFFF"))
    c.setFont("Helvetica-Bold", 24)
    c.drawString(59, H - 140, "Card heading")
    c.setFillColor(HexColor("#B11F25"))
    c.rect(58, H - 164, 26, 4, stroke=0, fill=1)
    body_lines(59, 190, 12, "#FFFFFF")
    body_lines(318, 190, 12, "#FFFFFF")
    c.setStrokeColor(HexColor("#FFFFFF"))
    c.setLineWidth(0.25)
    c.line(297, H - 185, 297, H - 410)
    c.setFillColor(HexColor("#FFFFFF"))
    c.setFont("Helvetica-Bold", 24)
    c.drawString(36, H - 540, "Second heading")
    c.setFillColor(HexColor("#ED4E46"))
    c.rect(36, H - 581, 26, 4, stroke=0, fill=1)
    body_lines(36, 610, 10, "#FFFFFF")
    body_lines(318, 610, 10, "#FFFFFF")
    c.line(297, H - 606, 297, H - 770)
    footer(2, "#FFFFFF")
    c.showPage()
    # 3 body: light page, two columns
    c.setFillColor(HexColor("#F0F0EE"))
    c.rect(0, 0, W, H, stroke=0, fill=1)
    for x in (36, 317):
        c.setFillColor(HexColor("#000000"))
        c.setFont("Helvetica-Bold", 24)
        c.drawString(x, H - 120, "Text heading")
        c.setFillColor(HexColor("#ED4E46"))
        c.rect(x, H - 145, 26, 4, stroke=0, fill=1)
        body_lines(x, 175, 30, "#000000")
    c.setStrokeColor(HexColor("#000000"))
    c.line(297, H - 90, 297, H - 778)
    footer(3, "#000000")
    c.showPage()
    # 4 back cover
    c.setFillColor(HexColor("#000000"))
    c.rect(0, 0, W, H, stroke=0, fill=1)
    c.setFillColor(HexColor("#FFFFFF"))
    c.setFont("Helvetica", 13)
    c.drawString(36, H - 62, "Authored by")
    c.setFont("Helvetica-Bold", 13)
    c.drawString(36, H - 89, "Jane Doe")
    c.setFont("Helvetica", 13)
    c.drawString(36, H - 106, "Practice Head")
    c.setFont("Helvetica", 9.2)
    c.drawString(36, H - 690, "Test Company builds things for customers in many countries.")
    c.drawString(36, H - 703, "For more information, please contact: info@example.com")
    c.showPage()
    c.save()
    return path


@pytest.fixture(scope="module")
def template(tmp_path_factory):
    from backend.branddocs import service

    ref = _reference(tmp_path_factory.mktemp("bd") / "reference.pdf")
    return service.create_template("reference.pdf", ref, "Test template", "")


def test_template_learned_from_reference(template):
    from backend.branddocs import service

    t = service.load_template(template["id"])
    assert set(t.pages) >= {"cover", "intro", "body", "back"}
    assert t.pages["intro"].card and t.pages["intro"].card.color.upper() == "#ED4E46"
    assert t.pages["intro"].background.upper() == "#201B5A"
    assert len(t.pages["body"].columns) == 2 and t.pages["body"].divider
    assert t.accent.upper() == "#ED4E46" and t.dash.get("w", 0) > 20
    assert "{page}" in template["footer"] and "{year}" in template["footer"]
    assert t.cover.label_text == "White Paper" and t.cover.marker is not None
    assert t.back.author_label == "Authored by"
    assert any("info@example.com" in b for b in t.back.boilerplate)
    assert template["page_images"], "reference page previews"


def _convert(template_id: str, filename: str, src, text: str = "", **opts):
    from backend.branddocs import service

    m = service.create_doc(filename, src, text, template_id, {"generate_images": False, "use_llm": False, **opts})
    service._running[m["id"]].join(timeout=240)
    return service.load_meta(m["id"])


def test_convert_text_into_template(template):
    from backend.branddocs import service

    text = "# Data is just raw material\n\n" + "\n\n".join(
        f"## Section {k}\n\n{LOREM * 3}\n\n- First point for section {k}\n- Second point\n" for k in range(1, 7)) + \
        "\n\n## Conclusion\n\nAt the end of the day, people make the decisions.\n"
    m = _convert(template["id"], "notes.md", None, text, images=1, label="Point of View")
    assert m["status"] == "done", m.get("error")
    r = m["report"]
    assert r["pages"] >= 4
    assert r["fidelity"]["coverage"] >= 0.98, r["fidelity"]
    assert m["title"] == "Data is just raw material"
    assert any(i["key"] == "cover" for i in m["images"]) and not any(i["generated"] for i in m["images"])
    import fitz

    doc = fitz.open(str(service.doc_file(m["id"], m["pdf"])))
    first, last = " ".join(doc[0].get_text().split()), " ".join(doc[-1].get_text().split())
    assert "Data is just raw material" in first and "Point of View" in first
    assert "info@example.com" in last
    assert re.search(r"Page\s+2", doc[1].get_text())


def test_convert_word_document(template, tmp_path):
    from docx import Document

    d = Document()
    d.add_heading("Claims modernisation", 0)
    for k in range(1, 5):
        d.add_heading(f"Part {k}", 1)
        d.add_paragraph(LOREM * 2)
    t = d.add_table(rows=2, cols=2)
    t.cell(0, 0).text, t.cell(0, 1).text, t.cell(1, 0).text, t.cell(1, 1).text = "Measure", "After", "Cycle time", "6 days"
    d.add_heading("Conclusion", 1)
    d.add_paragraph("Start small and prove the gain.")
    p = tmp_path / "claims.docx"
    d.save(p)
    m = _convert(template["id"], "claims.docx", p)
    assert m["status"] == "done", m.get("error")
    assert m["report"]["conclusion"] == "Conclusion"
    assert m["report"]["fidelity"]["coverage"] >= 0.98


def test_author_and_heading_rules():
    from backend.branddocs.content import build
    from backend.papers.model import Inline
    from backend.papers.structure import Elem

    el = [Elem(kind="heading", runs=[Inline(t="How to think")], level=1), Elem(kind="heading", runs=[Inline(t="about AI in")], level=1),
          Elem(kind="heading", runs=[Inline(t="simple words?")], level=1), Elem(kind="para", runs=[Inline(t="Forget buzzwords.")]),
          Elem(kind="heading", runs=[Inline(t="Authored by")], level=2),
          Elem(kind="heading", runs=[Inline(t="Daniel Gomez Practice Head, DE&A - Core daniel.gomez@zensar.com")], level=2)]
    doc = build(el, "Title", "x.pdf")
    assert doc.sections[0].title == "How to think about AI in simple words?"
    assert doc.authors[0].name == "Daniel Gomez" and doc.authors[0].email == "daniel.gomez@zensar.com"
    assert doc.authors[0].detail == ["Practice Head, DE&A - Core"]


def test_image_prompts_and_grading():
    from PIL import Image

    from backend.branddocs import imagegen
    from backend.branddocs.model import ImageStyle

    p = imagegen.section_prompt("Why AI feels so overwhelming?", "People hear that AI will automate everything. Fear and pressure.", "dark style")
    assert "?" not in p and p.endswith("dark style") and "overwhelming" in p
    assert "nude" not in imagegen._safe("a nude scene")
    st = ImageStyle(lab_mean=[30, 20, -25], lab_std=[25, 15, 25])
    import numpy as np

    from backend.branddocs.reference import _rgb_to_lab

    src = Image.new("RGB", (32, 32), (200, 120, 40))
    out = imagegen.grade(src, st)
    dist = lambda im: float(np.linalg.norm(_rgb_to_lab(np.asarray(im).astype(float)).reshape(-1, 3).mean(0) - np.array(st.lab_mean)))
    assert dist(out) < 0.5 * dist(src)  # pulled towards the reference photos' colour
    art = imagegen.brand_art((64, 80), ["#201B5A", "#ED4E46", "#F0F0EE"])
    assert art.size == (64, 80)


def test_fit_to_reference_length_without_llm():
    from backend.branddocs.content import DocContent, DocSection
    from backend.branddocs.fit import clean_heading, fit, regroup, select_sentences
    from backend.papers.model import Block

    def sec(title, n, level=1):
        text = " ".join(f"Sentence {k} explains how retrieval improves answers with fresh company data in practice." for k in range(n))
        return DocSection(title=title, level=level, blocks=[Block(kind="para", runs=[Inline(t=text)])])

    from backend.papers.model import Inline
    secs = [sec("I. INTRODUCTION", 40), sec("A. Background", 30, 2), sec("II. METHODS", 60), sec("A. Data", 20, 2),
            sec("B. Models", 20, 2), sec("III. RESULTS", 50), sec("References", 30)]
    secs[-1].kind = "references"
    doc = DocContent(title="A long paper", sections=secs, conclusion=sec("Conclusion", 20))
    assert clean_heading("II. OVERVIEW OF RAG") == "Overview of RAG"
    parts = regroup(doc, 4)
    assert len(parts) == 4 and not any("References" in t for p in parts for t in p.titles)
    picked = select_sentences(parts[0].text, 60)
    assert 30 <= sum(len(x.split()) for x in picked) <= 90
    r = fit(doc, 400, 4, 60, use_llm=False)
    assert r.mode == "selected" and 250 <= r.words <= 560, r.words
    assert len(r.doc.sections) == 4 and r.doc.conclusion is not None
    kept = fit(doc, 0, 4, 60, use_llm=False)
    assert kept.mode == "kept"


def test_title_lines_break_after_hyphens_only():
    from backend.branddocs.compose import title_lines

    lines = title_lines("Retrieval-Augmented Generation for Large Language Models: A Survey", "Helvetica-Bold", 40, 330)
    assert lines and all("Augment" not in ln or "Augmented" in ln for ln in lines)
    assert " ".join(lines).replace("- ", "-").replace("-", "-") .count("Retrieval-") == 1
    assert title_lines("Supercalifragilisticexpialidocious", "Helvetica-Bold", 60, 200) is None


def test_length_option_in_conversion(template):
    text = "# Short note\n\n" + "\n\n".join(f"## Point {k}\n\n{LOREM * 6}" for k in range(1, 9)) + "\n\n## Conclusion\n\nKeep it simple.\n"
    m = _convert(template["id"], "long.md", None, text, length="custom", words=300)
    assert m["status"] == "done", m.get("error")
    ln = m["report"]["length"]
    assert ln["mode"] == "selected" and ln["source_words"] > 1000 and ln["words"] < 600
    full = _convert(template["id"], "long.md", None, text, length="full")
    assert full["report"]["length"]["mode"] == "full" and full["report"]["words"] > 1000


def test_trim_removes_closing_sentences_only():
    from backend.branddocs.content import DocSection
    from backend.branddocs.fit import trim
    from backend.papers.model import Block, Inline

    para = " ".join(f"Point {k} keeps the original wording of the source." for k in range(12))
    secs = [DocSection(title=f"S{i}", blocks=[Block(kind="para", runs=[Inline(t=para)])]) for i in range(3)]
    before = {s.text[:40] for s in secs}
    trim(secs, [60, 60, 60], 200)
    assert sum(s.words for s in secs) <= 200 * 1.04
    assert {s.text[:40] for s in secs} == before  # openings untouched, nothing rewritten
    assert all(s.text.endswith(".") for s in secs)
