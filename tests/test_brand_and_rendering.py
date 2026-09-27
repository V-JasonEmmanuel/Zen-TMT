"""Brand engine (no invented values, provenance, analysis) and PPTX/PNG/SVG rendering."""
from pathlib import Path

import pytest
from PIL import Image
from pptx import Presentation

from backend.branding.brand_profile import BrandProfile, BrandStore, apply_value, ensure_default_brands
from backend.branding.reference_analyzer import analyze_references
from backend.branding.template_parser import analyze_pptx
from backend.branding.theme import NEUTRAL, build_theme
from backend.rendering.images.raster import RasterRenderer
from backend.rendering.images.svg import SVGRenderer
from backend.rendering.layouts import RenderContext, build_slide_scene, build_visual_scene
from backend.rendering.ppt.generator import PPTXGenerator


def test_zensar_profile_starts_empty_and_reports_fallbacks():
    ensure_default_brands()
    b = BrandStore().load("zensar")
    assert b.colors.primary is None and b.typography.heading.family is None  # nothing assumed
    assert b.status == "unconfigured"
    t = build_theme(b)
    assert "colors.primary" in t.fallbacks and "typography.heading.family" in t.fallbacks
    assert t.colors["primary"] == NEUTRAL["primary"]  # neutral, clearly-labelled placeholder
    assert not t.is_branded


def test_apply_value_validates_and_records_provenance():
    b = BrandProfile(id="demo", name="Demo")
    b = apply_value(b, "colors.primary", "#1a2b3c", "pptx_template", 0.6, "Accent 1")
    assert b.colors.primary == "#1A2B3C"
    assert b.provenance["colors.primary"].source == "pptx_template"
    with pytest.raises(Exception):
        apply_value(b, "colors.primary", "blue-ish")
    with pytest.raises(ValueError):
        apply_value(b, "id", "hack")
    b = apply_value(b, "typography.sizes.body", 15)
    assert b.typography.sizes["body"] == 15


def _demo_theme(tmp_path):
    store = BrandStore(tmp_path / "brands")
    b = store.create("Demo Brand")
    for path, val in (("colors.primary", "#0B3D91"), ("colors.accent", "#E4572E"), ("colors.text_primary", "#111111"),
                      ("colors.secondary", "#3A6EA5"), ("colors.surface", "#EEF2F7"), ("colors.background", "#FFFFFF"),
                      ("colors.text_secondary", "#555555"), ("typography.heading.family", "Arial"), ("typography.body.family", "Arial")):
        b = apply_value(b, path, val)
    store.save(b)
    return build_theme(b, store)


def test_configured_theme_uses_brand_tokens(tmp_path):
    t = _demo_theme(tmp_path)
    assert t.fallbacks == []
    assert t.colors["primary"] == "#0B3D91" and t.c("accent") == "#E4572E"
    assert t.c("on_primary") == "#FFFFFF"  # contrast-derived text on the brand primary


def test_pptx_output_opens_with_expected_content(extractive_plan, tmp_path):
    theme = _demo_theme(tmp_path)
    ctx = RenderContext(extractive_plan)
    scenes = [build_slide_scene(s, theme, ctx) for s in extractive_plan.slides]
    out = PPTXGenerator(theme).build(scenes, tmp_path / "deck.pptx", extractive_plan.title)
    prs = Presentation(str(out))
    assert len(prs.slides) == len(extractive_plan.slides)
    first_text = " ".join(sh.text_frame.text for sh in prs.slides[0].shapes if sh.has_text_frame)
    assert "Adaptive Retrieval" in first_text
    chart_slides = [s for s in prs.slides if any(sh.has_chart for sh in s.shapes)]
    assert chart_slides, "charts must be native, editable PowerPoint charts"
    # sources are carried into speaker notes, and brand colours are applied
    notes = [s.notes_slide.notes_text_frame.text for s in prs.slides if s.has_notes_slide]
    assert any("Sources:" in n and "sample_research.pdf" in n for n in notes)
    fills = {str(sh.fill.fore_color.rgb) for s in prs.slides for sh in s.shapes
             if getattr(sh, "fill", None) is not None and sh.shape_type == 1 and sh.fill.type == 1}
    assert "E4572E" in fills or "0B3D91" in fills
    unexpected = fills - {"0B3D91", "E4572E", "111111", "3A6EA5", "EEF2F7", "FFFFFF", "555555"}
    assert not unexpected, f"non-brand colours used: {unexpected}"


def test_raster_and_svg_visuals(extractive_plan, tmp_path):
    theme = _demo_theme(tmp_path)
    ctx = RenderContext(extractive_plan)
    rr, svg = RasterRenderer(960), SVGRenderer()
    made = 0
    for s in extractive_plan.slides:
        v = build_visual_scene(s, theme, ctx)
        if v is None:
            continue
        p = rr.save(v, tmp_path / f"{v.name}.png")
        with Image.open(p) as im:
            assert im.size == (960, 540)
        text = svg.render(v)
        assert text.startswith("<svg") and "</svg>" in text
        made += 1
    assert made >= 3


def test_template_analysis_suggests_but_does_not_apply(tmp_path):
    prs = Presentation()
    p = tmp_path / "tpl.pptx"
    prs.save(p)
    a = analyze_pptx(p)
    paths = {s["path"] for s in a["suggestions"]}
    assert {"typography.heading.family", "layout.slide_width_in", "colors.text_primary"} <= paths
    assert all(0 < s["confidence"] <= 1 for s in a["suggestions"])
    assert a["slide_size"]["width_in"] == 10.0


def test_reference_image_analysis(tmp_path):
    img = Image.new("RGB", (800, 450), "#FFFFFF")
    img.paste(Image.new("RGB", (800, 90), "#0B3D91"), (0, 0))
    img.paste(Image.new("RGB", (200, 40), "#E4572E"), (60, 200))
    p = tmp_path / "ref.png"
    img.save(p)
    r = analyze_references([p])
    sug = {s["path"]: s for s in r["suggestions"]}
    assert sug["colors.background"]["value"] in ("#FFFFFF", "#FCFCFC", "#FEFEFE")
    assert sug["colors.primary"]["confidence"] < 0.5  # pixel guesses are always low confidence
    assert r["images"][0]["full_width_bands"] >= 1
    assert Path(p).exists()
